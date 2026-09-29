import csv
import io
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, A4
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm


def exportar_relatorio(titulo, colunas, linhas, formato, nome_arquivo):
    """
    Gera a resposta HTTP com o arquivo no formato pedido.
    - titulo: título exibido no PDF
    - colunas: lista de nomes das colunas (cabeçalho)
    - linhas: lista de listas, cada uma é uma linha de dados (já em string/formatada)
    - formato: 'csv', 'xlsx' ou 'pdf'
    - nome_arquivo: nome do arquivo sem extensão
    """
    if formato == 'xlsx':
        return _exportar_xlsx(colunas, linhas, nome_arquivo)
    if formato == 'pdf':
        return _exportar_pdf(titulo, colunas, linhas, nome_arquivo)
    return _exportar_csv(colunas, linhas, nome_arquivo)


def _exportar_csv(colunas, linhas, nome_arquivo):
    response = HttpResponse(content_type='text/csv; charset=utf-8-sig')
    response['Content-Disposition'] = f'attachment; filename="{nome_arquivo}.csv"'
    writer = csv.writer(response, delimiter=';')
    writer.writerow(colunas)
    writer.writerows(linhas)
    return response


def _exportar_xlsx(colunas, linhas, nome_arquivo):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Relatório'

    cabecalho_fill = PatternFill(start_color='7C5CFF', end_color='7C5CFF', fill_type='solid')
    cabecalho_font = Font(color='FFFFFF', bold=True)

    ws.append(colunas)
    for cell in ws[1]:
        cell.fill = cabecalho_fill
        cell.font = cabecalho_font

    for linha in linhas:
        ws.append(linha)

    for coluna in ws.columns:
        maior = max((len(str(c.value)) for c in coluna if c.value is not None), default=10)
        ws.column_dimensions[coluna[0].column_letter].width = min(maior + 3, 40)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{nome_arquivo}.xlsx"'
    return response


def _exportar_pdf(titulo, colunas, linhas, nome_arquivo):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        topMargin=1.5 * cm,
        bottomMargin=1.5 * cm,
        leftMargin=1.5 * cm,
        rightMargin=1.5 * cm,
    )
    estilos = getSampleStyleSheet()
    elementos = [Paragraph(titulo, estilos['Title']), Spacer(1, 0.5 * cm)]

    dados_tabela = [colunas] + linhas
    tabela = Table(dados_tabela, repeatRows=1)
    tabela.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#7C5CFF')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F0F0F5')]),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    elementos.append(tabela)

    doc.build(elementos)
    buffer.seek(0)

    response = HttpResponse(buffer.read(), content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{nome_arquivo}.pdf"'
    return response